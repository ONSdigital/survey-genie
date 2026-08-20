#!/usr/bin/env python3
"""Convert a Survey Genie Excel workbook into a validated JSON definition."""

from __future__ import annotations

import argparse
import json
import logging
import os
from pathlib import Path
import sys
import tempfile
from typing import cast

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from survey_genie.survey.loader import SurveyDefinitionError, load_survey_definition

logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
logger = logging.getLogger(__name__)

MAX_WORKBOOK_BYTES = 10 * 1024 * 1024
REQUIRED_SHEETS = {"Survey", "Pages", "Options"}
VALID_SECTIONS = {"survey", "feedback"}
VALID_PAGE_TYPES = {"question", "guidance"}
VALID_ANSWER_TYPES = {"radio", "text"}


class SurveyWorkbookError(ValueError):
    """Raised when a survey workbook cannot be converted safely."""


def main() -> int:
    """Convert a Survey Genie Excel workbook to JSON.

    Returns:
        int: Zero when conversion succeeds.
    """
    parser = argparse.ArgumentParser(
        description="Convert a Survey Genie XLSX definition to validated JSON."
    )
    parser.add_argument("input", type=Path, help="Path to the source .xlsx workbook.")
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("survey_definition.json"),
        help="Output JSON path. Defaults to survey_definition.json.",
    )
    args = parser.parse_args()

    definition = convert_workbook(args.input)
    write_validated_definition(definition, args.output)

    logger.info("Wrote validated survey definition to %s", args.output)
    return 0


def convert_workbook(workbook_path: Path) -> dict[str, object]:
    """Convert workbook content into a Survey Genie definition.

    Args:
        workbook_path: Path to an XLSX workbook following the authoring schema.

    Returns:
        dict[str, object]: Survey Genie JSON-compatible definition.

    Raises:
        SurveyWorkbookError: If the workbook is missing, malformed, too large,
            uses unsupported values, or violates workbook-level rules.
    """
    # pylint: disable=too-many-locals,too-many-branches
    _validate_workbook_path(workbook_path)

    try:
        workbook = load_workbook(
            workbook_path,
            read_only=True,
            data_only=True,
        )
    except (InvalidFileException, OSError, ValueError) as exc:
        raise SurveyWorkbookError(f"Unable to read workbook {workbook_path}: {exc}") from exc

    try:
        missing_sheets = REQUIRED_SHEETS.difference(workbook.sheetnames)
        if missing_sheets:
            missing = ", ".join(sorted(missing_sheets))
            raise SurveyWorkbookError(f"Workbook is missing required sheet(s): {missing}")

        survey_rows = _sheet_rows(workbook["Survey"])
        page_rows = _sheet_rows(workbook["Pages"])
        option_rows = _sheet_rows(workbook["Options"])

        if len(survey_rows) != 1:
            raise SurveyWorkbookError("Survey sheet must contain exactly one populated data row")

        survey = survey_rows[0]
        schema_version = _required_positive_int(survey, "schema_version", "Survey")
        survey_title = _required_string(survey, "survey_title", "Survey")
        wave_id = _required_string(survey, "wave_id", "Survey")
        intro_enabled = _required_bool(survey, "intro_enabled", "Survey")
        survey_enabled = _required_bool(survey, "survey_enabled", "Survey")
        feedback_enabled = _required_bool(survey, "feedback_enabled", "Survey")

        if schema_version != 1:
            raise SurveyWorkbookError(f"Survey.schema_version must be 1, got {schema_version}")

        if not survey_enabled:
            raise SurveyWorkbookError("Survey Genie requires the survey section to be enabled")

        if intro_enabled:
            raise SurveyWorkbookError(
                "Intro authoring is reserved in the workbook schema but is not "
                "implemented by this MVP converter"
            )

        survey_start_page_id = _required_string(
            survey,
            "survey_start_page_id",
            "Survey",
        )
        feedback_start_page_id = _optional_string(
            survey,
            "feedback_start_page_id",
        )

        survey_pages = _build_section_pages(
            page_rows,
            option_rows,
            section="survey",
        )
        feedback_pages = _build_section_pages(
            page_rows,
            option_rows,
            section="feedback",
        )

        if not survey_pages:
            raise SurveyWorkbookError("Pages sheet must contain at least one survey page")

        if feedback_enabled:
            if not feedback_start_page_id:
                raise SurveyWorkbookError(
                    "Survey.feedback_start_page_id is required when feedback is enabled"
                )
            if not feedback_pages:
                raise SurveyWorkbookError(
                    "Pages sheet must contain feedback pages when feedback is enabled"
                )
        elif feedback_pages:
            raise SurveyWorkbookError(
                "Pages sheet contains feedback rows but Survey.feedback_enabled is FALSE"
            )

        definition: dict[str, object] = {
            "schema_version": schema_version,
            "survey_title": survey_title,
            "wave_id": wave_id,
            "survey_intro": {"enabled": False},
            "survey_pages": {
                "enabled": True,
                "start_page_id": survey_start_page_id,
                "pages": survey_pages,
            },
            "survey_feedback": {
                "enabled": feedback_enabled,
            },
        }

        if feedback_enabled:
            feedback = cast(dict[str, object], definition["survey_feedback"])
            feedback["start_page_id"] = feedback_start_page_id
            feedback["pages"] = feedback_pages

        return definition
    finally:
        workbook.close()


def write_validated_definition(
    definition: dict[str, object],
    output_path: Path,
) -> None:
    """Validate and atomically write a generated survey definition.

    The generated JSON is first written to a temporary file and passed through
    the existing Survey Genie loader. The destination is replaced only if the
    generated definition is valid.

    Args:
        definition: Generated Survey Genie definition.
        output_path: Destination JSON file.

    Raises:
        SurveyWorkbookError: If the JSON cannot be written or fails Survey
            Genie validation.
    """
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None

    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            suffix=".json",
            prefix=".survey-genie-",
            dir=output_path.parent,
            delete=False,
        ) as temporary_file:
            json.dump(definition, temporary_file, indent=2)
            temporary_file.write("\n")
            temporary_path = Path(temporary_file.name)

        load_survey_definition(temporary_path)
        os.replace(temporary_path, output_path)
        temporary_path = None
    except SurveyDefinitionError as exc:
        raise SurveyWorkbookError(
            f"Generated survey definition failed Survey Genie validation: {exc}"
        ) from exc
    except OSError as exc:
        raise SurveyWorkbookError(
            f"Unable to write survey definition to {output_path}: {exc}"
        ) from exc
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)


def _validate_workbook_path(workbook_path: Path) -> None:
    """Validate the workbook before parsing it.

    Args:
        workbook_path: Workbook path to validate.

    Raises:
        SurveyWorkbookError: If the path is invalid or the file is too large.
    """
    if workbook_path.suffix.lower() != ".xlsx":
        raise SurveyWorkbookError("Input file must use the .xlsx extension")

    try:
        file_size = workbook_path.stat().st_size
    except FileNotFoundError as exc:
        raise SurveyWorkbookError(f"Workbook was not found: {workbook_path}") from exc
    except OSError as exc:
        raise SurveyWorkbookError(f"Unable to inspect workbook {workbook_path}: {exc}") from exc

    if file_size > MAX_WORKBOOK_BYTES:
        raise SurveyWorkbookError(
            f"Workbook exceeds the {MAX_WORKBOOK_BYTES // (1024 * 1024)} MB limit"
        )


def _sheet_rows(worksheet: Worksheet) -> list[dict[str, object]]:
    """Read a worksheet as dictionaries keyed by its header row.

    Args:
        worksheet: Worksheet to read.

    Returns:
        list[dict[str, object]]: Populated data rows.

    Raises:
        SurveyWorkbookError: If headers are blank or duplicated.
    """
    raw_rows = list(worksheet.iter_rows(values_only=True))

    if not raw_rows:
        raise SurveyWorkbookError(f"{worksheet.title} sheet is empty")

    raw_headers = raw_rows[0]
    headers: list[str] = []

    for column_index, value in enumerate(raw_headers, start=1):
        if value is None or not str(value).strip():
            raise SurveyWorkbookError(f"{worksheet.title} header column {column_index} is blank")
        headers.append(str(value).strip())

    if len(headers) != len(set(headers)):
        raise SurveyWorkbookError(f"{worksheet.title} sheet contains duplicate column headers")

    rows: list[dict[str, object]] = []
    for row_number, values in enumerate(raw_rows[1:], start=2):
        if not any(value not in (None, "") for value in values):
            continue

        padded_values = list(values) + [None] * (len(headers) - len(values))
        row = dict(zip(headers, padded_values, strict=False))
        row["_row_number"] = row_number
        rows.append(row)

    return rows


def _build_section_pages(
    page_rows: list[dict[str, object]],
    option_rows: list[dict[str, object]],
    *,
    section: str,
) -> list[dict[str, object]]:
    """Build ordered pages for one survey section.

    Args:
        page_rows: All workbook page rows.
        option_rows: All workbook radio option rows.
        section: Section to build.

    Returns:
        list[dict[str, object]]: Ordered Survey Genie pages.

    Raises:
        SurveyWorkbookError: If page or option data is invalid.
    """
    if section not in VALID_SECTIONS:
        raise SurveyWorkbookError(f"Unsupported section: {section}")

    rows = [row for row in page_rows if _optional_string(row, "section") == section]
    rows.sort(key=lambda row: _required_positive_int(row, "sequence", "Pages"))

    sequences: set[int] = set()
    page_ids: set[str] = set()
    pages: list[dict[str, object]] = []

    for row in rows:
        sequence = _required_positive_int(row, "sequence", "Pages")
        page_id = _required_string(row, "page_id", "Pages")

        if sequence in sequences:
            raise SurveyWorkbookError(
                f"Pages contains duplicate sequence {sequence} in section {section!r}"
            )
        if page_id in page_ids:
            raise SurveyWorkbookError(
                f"Pages contains duplicate page_id {page_id!r} in section {section!r}"
            )

        sequences.add(sequence)
        page_ids.add(page_id)

        page_type = _required_string(row, "page_type", "Pages")
        if page_type not in VALID_PAGE_TYPES:
            raise _row_error(
                row,
                "Pages",
                f"page_type must be one of {sorted(VALID_PAGE_TYPES)}",
            )

        if section == "feedback" and page_type != "question":
            raise _row_error(
                row,
                "Pages",
                "current Survey Genie feedback pages must have page_type 'question'",
            )

        if page_type == "question":
            pages.append(
                _build_question_page(
                    row,
                    option_rows,
                    section=section,
                    page_id=page_id,
                )
            )
        else:
            pages.append(_build_guidance_page(row, page_id=page_id))

    return pages


def _build_question_page(
    row: dict[str, object],
    option_rows: list[dict[str, object]],
    *,
    section: str,
    page_id: str,
) -> dict[str, object]:
    """Build a question page.

    Args:
        row: Source Pages row.
        option_rows: All workbook option rows.
        section: Page section.
        page_id: Page identifier.

    Returns:
        dict[str, object]: Survey Genie question page.

    Raises:
        SurveyWorkbookError: If question configuration is invalid.
    """
    answer_type = _required_string(row, "answer_type", "Pages")
    if answer_type not in VALID_ANSWER_TYPES:
        raise _row_error(
            row,
            "Pages",
            f"answer_type must be one of {sorted(VALID_ANSWER_TYPES)}",
        )

    question: dict[str, object] = {"text": _required_string(row, "question_text", "Pages")}
    description = _optional_string(row, "question_description")
    if description:
        question["description"] = description

    answer: dict[str, object] = {
        "type": answer_type,
        "name": _required_string(row, "answer_name", "Pages"),
        "required": _required_bool(row, "required", "Pages"),
    }

    answer_label = _optional_string(row, "answer_label")
    if answer_label:
        answer["label"] = answer_label

    if answer_type == "radio":
        answer["options"] = _build_radio_options(
            option_rows,
            section=section,
            page_id=page_id,
        )
    else:
        multiline = _optional_bool(row, "multiline", "Pages")
        if multiline is not None:
            answer["multiline"] = multiline

        rows = _optional_positive_int(row, "rows", "Pages")
        if rows is not None:
            answer["rows"] = rows

        character_limit = _optional_positive_int(
            row,
            "character_limit",
            "Pages",
        )
        if character_limit is not None:
            answer["character_limit"] = character_limit

        placeholder = _optional_string(row, "placeholder")
        if placeholder is not None:
            answer["placeholder"] = placeholder

    return {
        "page_id": page_id,
        "page_type": "question",
        "page_title": _required_string(row, "page_title", "Pages"),
        "question_name": _required_string(row, "question_name", "Pages"),
        "question": question,
        "answer": answer,
        "submit_button": {"text": _required_string(row, "button_text", "Pages")},
    }


def _build_guidance_page(
    row: dict[str, object],
    *,
    page_id: str,
) -> dict[str, object]:
    """Build a guidance page.

    Args:
        row: Source Pages row.
        page_id: Page identifier.

    Returns:
        dict[str, object]: Survey Genie guidance page.
    """
    return {
        "page_id": page_id,
        "page_type": "guidance",
        "page_title": _required_string(row, "page_title", "Pages"),
        "guidance_overview": _required_string(
            row,
            "guidance_overview",
            "Pages",
        ),
        "guidance_subsection": _required_string(
            row,
            "guidance_subsection",
            "Pages",
        ),
        "continue_button": {"text": _required_string(row, "button_text", "Pages")},
    }


def _build_radio_options(
    option_rows: list[dict[str, object]],
    *,
    section: str,
    page_id: str,
) -> list[dict[str, object]]:
    """Build ordered radio options for a page.

    Args:
        option_rows: All workbook option rows.
        section: Parent section.
        page_id: Parent page identifier.

    Returns:
        list[dict[str, object]]: Ordered radio option definitions.

    Raises:
        SurveyWorkbookError: If no options exist or option order is duplicated.
    """
    rows = [
        row
        for row in option_rows
        if _optional_string(row, "section") == section
        and _optional_string(row, "page_id") == page_id
    ]
    rows.sort(key=lambda row: _required_positive_int(row, "sequence", "Options"))

    if not rows:
        raise SurveyWorkbookError(
            f"Radio page {page_id!r} in section {section!r} has no Options rows"
        )

    sequences: set[int] = set()
    options: list[dict[str, object]] = []

    for row in rows:
        sequence = _required_positive_int(row, "sequence", "Options")
        if sequence in sequences:
            raise _row_error(
                row,
                "Options",
                f"duplicate sequence {sequence} for page {page_id!r}",
            )
        sequences.add(sequence)

        option: dict[str, object] = {
            "id": _required_string(row, "option_id", "Options"),
            "label": _required_string(row, "label", "Options"),
            "value": _required_string(row, "value", "Options"),
        }

        target_page_id = _optional_string(row, "target_page_id")
        if target_page_id:
            option["target_page_id"] = target_page_id

        options.append(option)

    return options


def _required_string(
    row: dict[str, object],
    field_name: str,
    sheet_name: str,
) -> str:
    """Return a required non-empty string value."""
    value = _optional_string(row, field_name)
    if not value:
        raise _row_error(row, sheet_name, f"{field_name} is required")
    return value


def _optional_string(
    row: dict[str, object],
    field_name: str,
) -> str | None:
    """Return a stripped string or None for an empty cell."""
    value = row.get(field_name)
    if value is None:
        return None

    text = str(value).strip()
    return text or None


def _required_bool(
    row: dict[str, object],
    field_name: str,
    sheet_name: str,
) -> bool:
    """Return a required boolean cell value."""
    value = _optional_bool(row, field_name, sheet_name)
    if value is None:
        raise _row_error(row, sheet_name, f"{field_name} is required")
    return value


def _optional_bool(
    row: dict[str, object],
    field_name: str,
    sheet_name: str,
) -> bool | None:
    """Return an optional boolean from an Excel boolean or TRUE/FALSE text."""
    value = row.get(field_name)
    if value in (None, ""):
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, str):
        normalised = value.strip().lower()
        if normalised == "true":
            return True
        if normalised == "false":
            return False

    raise _row_error(
        row,
        sheet_name,
        f"{field_name} must be TRUE or FALSE",
    )


def _required_positive_int(
    row: dict[str, object],
    field_name: str,
    sheet_name: str,
) -> int:
    """Return a required positive integer cell value."""
    value = _optional_positive_int(row, field_name, sheet_name)
    if value is None:
        raise _row_error(row, sheet_name, f"{field_name} is required")
    return value


def _optional_positive_int(
    row: dict[str, object],
    field_name: str,
    sheet_name: str,
) -> int | None:
    """Return an optional positive integer cell value."""
    value = row.get(field_name)
    if value in (None, ""):
        return None

    if isinstance(value, bool):
        raise _row_error(
            row,
            sheet_name,
            f"{field_name} must be a positive integer",
        )

    if isinstance(value, int):
        parsed = value
    elif isinstance(value, float) and value.is_integer():
        parsed = int(value)
    else:
        raise _row_error(
            row,
            sheet_name,
            f"{field_name} must be a positive integer",
        )

    if parsed < 1:
        raise _row_error(
            row,
            sheet_name,
            f"{field_name} must be a positive integer",
        )

    return parsed


def _row_error(
    row: dict[str, object],
    sheet_name: str,
    message: str,
) -> SurveyWorkbookError:
    """Create a workbook error containing the source row number."""
    row_number = row.get("_row_number", "?")
    return SurveyWorkbookError(f"{sheet_name} row {row_number}: {message}")


if __name__ == "__main__":
    try:
        sys.exit(main())
    except SurveyWorkbookError as exc:
        logger.error("%s", exc)
        sys.exit(2)
